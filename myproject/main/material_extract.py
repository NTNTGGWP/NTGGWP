"""從課程教材檔案（pdf / docx / pptx / txt）抽取純文字，供 RAG 建索引使用。

檔案可能存在本機（開發）或 S3（正式）；統一用 Django 的 storage 開檔，
兩種環境都適用。無法抽取或格式不支援時回傳空字串（呼叫端會略過）。
"""
import io
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXT = {'pdf', 'docx', 'pptx', 'txt'}


def _ext(name):
    return name.rsplit('.', 1)[-1].lower() if '.' in name else ''


def extract_text(material):
    """讀取 LessonMaterial 的檔案並回傳純文字，失敗回傳空字串。"""
    name = material.file.name
    ext = _ext(name)
    if ext not in SUPPORTED_EXT:
        return ''

    try:
        with material.file.open('rb') as fh:
            data = fh.read()
    except Exception as exc:
        logger.warning('教材開檔失敗 %s: %s', name, exc)
        return ''

    try:
        if ext == 'txt':
            return _from_txt(data)
        if ext == 'pdf':
            return _from_pdf(data)
        if ext == 'docx':
            return _from_docx(data)
        if ext == 'pptx':
            return _from_pptx(data)
    except Exception as exc:
        logger.warning('教材抽取失敗 %s: %s', name, exc)
        return ''
    return ''


def _from_txt(data):
    for enc in ('utf-8', 'big5', 'cp950', 'latin-1'):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return ''


def _from_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return '\n'.join((page.extract_text() or '') for page in reader.pages)


def _from_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _from_pptx(data):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = ''.join(run.text for run in para.runs).strip()
                    if text:
                        slide_parts.append(text)
        if slide_parts:
            parts.append(f'【投影片 {i}】\n' + '\n'.join(slide_parts))
    return '\n\n'.join(parts)


def chunk_text(text, size=600, overlap=100):
    """把長文字切成有重疊的片段，方便 RAG 檢索與控制 token。"""
    text = ' '.join(text.split())
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        chunks.append(text[start:end])
        if end >= n:
            break
        start = end - overlap
    return chunks
