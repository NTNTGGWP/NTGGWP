"""把 export_powerbi_json 產生的 JSON 組成 Power BI 用的 .xlsx 工作簿。

這支取代原本依賴 OpenAI 專屬工具的 build_powerbi_workbook.mjs —— 只用 openpyxl，
你自己就能重建，不需要 GPT。產出的工作表名稱、欄位、命名表格都與既有的
LearnMate_AI_Dashboard.pbix 對齊，.pbix 直接「重新整理」即可套用新資料。

用法：
    python build_workbook.py --input ../outputs/xxx/learnmate_powerbi_data.json \
                             --output ../outputs/xxx/LearnMate_PowerBI_Data.xlsx

一步到位（先導出再建簿）：
    cd myproject
    python manage.py export_powerbi_json --output ../outputs/learnmate_powerbi_data.json
    cd ../powerbi
    python build_workbook.py --input ../outputs/learnmate_powerbi_data.json \
                             --output ../outputs/LearnMate_PowerBI_Data.xlsx
"""
import argparse
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# 17 張資料表的合約（與 .pbix 對齊，欄位順序固定）
SPECS = [
    ("Students", "students", "StudentsTable", ["student_id", "username", "display_name", "role", "is_active", "date_joined"]),
    ("Courses", "courses", "CoursesTable", ["course_id", "course_title", "teacher_id", "teacher_name", "category", "level", "is_published", "created_at"]),
    ("Chapters", "chapters", "ChaptersTable", ["chapter_id", "course_id", "chapter_title", "sort_order", "created_at"]),
    ("Lessons", "lessons", "LessonsTable", ["lesson_id", "chapter_id", "lesson_title", "duration_minutes", "sort_order", "is_free_preview", "created_at"]),
    ("Enrollments", "enrollments", "EnrollmentsTable", ["enrollment_id", "student_id", "course_id", "purchased_at"]),
    ("LessonProgress", "lesson_progress", "LessonProgressTable", ["progress_id", "student_id", "course_id", "lesson_id", "watched_seconds", "total_watch_seconds", "duration_seconds", "watch_rate", "view_count", "page_open_count", "replayed_seconds", "rewatch_ratio", "last_position_seconds", "is_completed", "updated_at"]),
    ("LearningRecords", "learning_records", "LearningRecordsTable", ["learning_record_id", "student_id", "course_id", "lesson_id", "minutes", "watched_at"]),
    ("CourseSummaries", "course_summaries", "CourseSummariesTable", ["summary_id", "course_id", "summary", "core_concept_count", "key_term_count", "created_at", "updated_at"]),
    ("Quizzes", "quizzes", "QuizzesTable", ["quiz_id", "course_id", "quiz_type", "source_result_id", "created_by_id", "created_at"]),
    ("Questions", "questions", "QuestionsTable", ["question_id", "quiz_id", "question_order", "question_text", "topic_tag", "correct_index", "explanation"]),
    ("QuizResults", "quiz_results", "QuizResultsTable", ["result_id", "student_id", "quiz_id", "course_id", "quiz_type", "source_result_id", "accuracy_rate", "score_percent", "attempt_number", "previous_score_percent", "score_change_percent", "is_latest_attempt", "weak_topic_count", "strong_topic_count", "overall_feedback", "submitted_at"]),
    ("Answers", "answers", "AnswersTable", ["answer_id", "result_id", "student_id", "course_id", "quiz_id", "question_id", "topic_tag", "selected_index", "correct_index", "is_correct"]),
    ("WeakTopics", "weak_topics", "WeakTopicsTable", ["result_id", "student_id", "course_id", "quiz_id", "topic", "issue_description", "severity", "submitted_at"]),
    ("RemedialCompare", "remedial_comparisons", "RemedialCompareTable", ["student_id", "course_id", "before_result_id", "after_result_id", "before_accuracy_rate", "after_accuracy_rate", "improvement_rate", "before_submitted_at", "after_submitted_at"]),
    ("StudentCourseMetrics", "student_course_metrics", "StudentCourseMetricsTable", ["student_id", "course_id", "lesson_count", "played_lesson_count", "completed_lesson_count", "lesson_completion_rate", "average_watch_rate", "quiz_attempt_count", "historical_average_score_percent", "latest_score_percent", "previous_score_percent", "latest_score_change_percent", "latest_quiz_type", "latest_weak_topic_count", "historical_weakness_mentions", "latest_submitted_at"]),
    ("CourseHealth", "course_health", "CourseHealthTable", ["course_id", "enrolled_students", "lesson_count", "recorded_progress_count", "completed_lesson_count", "lesson_completion_rate", "quiz_result_count", "average_quiz_accuracy_rate", "weakness_mentions"]),
    ("TopicStats", "topic_stats", "TopicStatsTable", ["course_id", "topic", "question_count", "answer_count", "wrong_answer_count", "error_rate", "weakness_mentions"]),
]

RATE_COLS = {"watch_rate", "rewatch_ratio", "accuracy_rate", "before_accuracy_rate",
             "after_accuracy_rate", "improvement_rate", "lesson_completion_rate",
             "average_watch_rate", "average_quiz_accuracy_rate", "error_rate"}
DATE_COLS_SUFFIX = ("_at",)
WRAP_COLS = {"summary", "question_text", "explanation", "overall_feedback", "issue_description"}

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="172033")


def _is_date_col(col):
    return col.endswith(DATE_COLS_SUFFIX) or col == "date_joined"


def _coerce(col, value):
    if value is None:
        return None
    if _is_date_col(col) and isinstance(value, str):
        try:
            dt = datetime.fromisoformat(value)
            return dt.replace(tzinfo=None)  # openpyxl 不接受 tz-aware
        except ValueError:
            return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _number_format(col):
    if _is_date_col(col):
        return "yyyy-mm-dd hh:mm"
    if col in RATE_COLS:
        return "0.0%"
    if col.endswith("_change_percent"):
        return "+0.0;-0.0;0.0"
    if col == "score_percent" or col.endswith("_score_percent"):
        return "0.0"
    if (col.endswith("_id") or col.endswith("_count") or "seconds" in col
            or col in ("minutes", "question_order", "sort_order")):
        return "#,##0"
    return None


def _width(col):
    if col in WRAP_COLS:
        return 48
    if _is_date_col(col):
        return 21
    if "title" in col or col in ("display_name", "teacher_name"):
        return 28
    if "topic" in col or col in ("category", "username"):
        return 22
    return max(16, min(27, len(col) + 2))


def add_sheet(wb, name, key, table_name, columns, rows):
    ws = wb.create_sheet(title=name)
    ws.sheet_view.showGridLines = False
    ws.append(columns)
    for row in rows:
        ws.append([_coerce(c, row.get(c)) for c in columns])

    n_rows = len(rows)
    last_col = get_column_letter(len(columns))

    # 標題列樣式
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # 欄寬、數字格式、換行
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = _width(col)
        fmt = _number_format(col)
        wrap = col in WRAP_COLS
        if fmt or wrap:
            for r in range(2, n_rows + 2):
                cell = ws.cell(row=r, column=i)
                if fmt:
                    cell.number_format = fmt
                if wrap:
                    cell.alignment = Alignment(wrap_text=True, vertical="center")

    # 命名表格（.pbix 以這些表格名取數）
    ref = f"A1:{last_col}{max(n_rows + 1, 2)}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    return n_rows


def build(input_path, output_path):
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)  # 移除預設工作表

    # 說明頁
    info = wb.create_sheet(title="說明")
    info.sheet_view.showGridLines = False
    info["A1"] = "LearnMate Power BI 資料包"
    info["A1"].font = Font(name="Arial", size=16, bold=True, color="0F172A")
    info["A2"] = f"來源：LearnMate 資料庫（export_powerbi_json）｜匯出時間：{data.get('generated_at', '')}｜僅含實際紀錄，未補造學習數據。"
    info["A2"].font = Font(name="Arial", size=10, italic=True, color="52606D")
    info["A4"] = "資料表"
    info["B4"] = "資料列數"
    for cell in (info["A4"], info["B4"]):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 12

    counts = []
    for (name, key, table_name, columns) in SPECS:
        rows = data.get(key, []) or []
        n = add_sheet(wb, name, key, table_name, columns, rows)
        counts.append((name, n))

    for idx, (name, n) in enumerate(counts, start=5):
        info[f"A{idx}"] = name
        info[f"B{idx}"] = n
    info_ref = f"A4:B{4 + len(counts)}"
    inv = Table(displayName="DataInventoryTable", ref=info_ref)
    inv.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    info.add_table(inv)

    # 把說明頁移到最前
    wb.move_sheet("說明", -(len(wb.sheetnames) - 1))
    wb.save(output_path)
    total = sum(n for _, n in counts)
    print(f"已建立工作簿：{output_path}")
    print(f"工作表 {len(counts) + 1} 張、資料列合計 {total}")
    for name, n in counts:
        print(f"  {name}={n}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="export_powerbi_json 產生的 JSON")
    ap.add_argument("--output", required=True, help="輸出的 .xlsx 路徑")
    args = ap.parse_args()
    build(args.input, args.output)
